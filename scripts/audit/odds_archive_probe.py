"""Probe free historical NFL odds archives (evidence only): Sportsbook Review Online and Australia Sports Betting.
Saves each page / file's first bytes and, for spreadsheets, sheet names, header rows and a few data rows.
Writes docs/evidence/odds_archives/<UTC stamp>/summary.json."""
from __future__ import annotations

import datetime as dt
import io
import json
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path

UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
PAGES = ["https://www.sportsbookreviewsonline.com/scoresoddsarchives/nfl/nfloddsarchives.htm",
         "https://sportsbookreviewsonline.com/scoresoddsarchives/nfl/nfloddsarchives.htm",
         "https://www.aussportsbetting.com/data/historical-nfl-results-and-odds-data/"]
FILES = ["https://www.aussportsbetting.com/historical_data/nfl.xlsx"]


def get(url):
    time.sleep(1)
    try:
        req = urllib.request.Request(url, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, r.headers.get("content-type", ""), r.read(20_000_000)
    except Exception as e:  # noqa: BLE001
        return str(e)[:150], "", b""


def sheet_summary(body):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True)
    out = {}
    for ws in wb.worksheets[:3]:
        rows = [[None if c is None else str(c) for c in r] for _, r in zip(range(6), ws.iter_rows(values_only=True))]
        out[ws.title] = {"first_rows": rows, "max_row": ws.max_row}
    return out


def main():
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    out = Path("docs/evidence/odds_archives") / stamp
    out.mkdir(parents=True, exist_ok=True)
    s = {"stamp": stamp, "pages": [], "files": []}
    links = []
    for u in PAGES:
        st, ct, b = get(u)
        t = b.decode("utf-8", "replace")
        found = [urllib.parse.urljoin(u, h) for h in re.findall(r'href="([^"]+\.(?:xlsx?|csv))"', t, re.I)]
        links += found
        s["pages"].append({"url": u, "status": st, "bytes": len(b), "spreadsheet_links": found[:60]})
        if b:
            (out / (re.sub(r"[^A-Za-z0-9]+", "_", u.split("//")[1])[:80] + ".html")).write_text(t[:60000])
        print(u, st, len(b), found[:10], flush=True)
    # SBRO links each season to an HTML page (/scoresoddsarchives/nfl-odds-YYYY-YY); follow two and look for tables / files
    for season_page in ["https://www.sportsbookreviewsonline.com/scoresoddsarchives/nfl-odds-2021-22",
                        "https://www.sportsbookreviewsonline.com/scoresoddsarchives/nfl-odds-2016-17"]:
        st, ct, b = get(season_page)
        t = b.decode("utf-8", "replace")
        files = [urllib.parse.urljoin(season_page, h) for h in re.findall(r'href="([^"]+\.(?:xlsx?|csv))"', t, re.I)]
        tables = re.findall(r"<table.*?</table>", t, re.S | re.I)
        rows = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " | ", r)).strip() for r in re.findall(r"<tr.*?</tr>", tables[0], re.S | re.I)[:8]] if tables else []
        s["pages"].append({"url": season_page, "status": st, "bytes": len(b), "spreadsheet_links": files[:20],
                           "n_tables": len(tables), "first_rows": rows, "n_rows_table0": len(re.findall(r"<tr", tables[0], re.I)) if tables else 0})
        (out / ("sbro_" + season_page.rsplit("/", 1)[1] + ".html")).write_text(t[:200000])
        print(season_page, st, len(b), files[:5], len(tables), rows[:4], flush=True)
        links += files
    sbro = [l for l in links if "sportsbookreviewsonline" in l]
    for u in FILES + sbro[:2]:
        st, ct, b = get(u)
        rec = {"url": u, "status": st, "content_type": ct, "bytes": len(b)}
        try:
            if b[:2] == b"PK":
                rec["sheets"] = sheet_summary(b)
            elif b:
                rec["head"] = b[:1500].decode("latin-1")
        except Exception as e:  # noqa: BLE001
            rec["parse_error"] = str(e)[:200]
        s["files"].append(rec)
        print(u, st, len(b), str(rec.get("sheets", rec.get("head", "")))[:600], flush=True)
    (out / "summary.json").write_text(json.dumps(s, indent=1, default=str))


if __name__ == "__main__":
    main()
